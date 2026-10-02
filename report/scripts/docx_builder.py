"""Minimal builder that re-uses the ORIGINAL docx (styles, numbering, footers, cover pages) and appends new content
using the same formatting conventions as the original (TH SarabunPSK 16 pt, Heading1/2, FigCaption, TabCaption, TableGrid)."""
import copy, re, os, sys, glob
from docx import Document
from docx.shared import Cm, Emu, Pt
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from lxml import etree
sys.path.insert(0, glob.glob("/root/.claude/skills/synced/*/thai-docx/scripts")[0])
sys.path.insert(0, glob.glob("/root/.claude/skills/synced/*/word-equation/scripts")[0])
from thai_docx import insert_zwsp
from word_equation import latex_to_omml

RED = "C00000"
HF = "TH SarabunPSK"
W = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
_bm = [2000]
def _el(tag, **attrs):
    e = OxmlElement(tag)
    for k, v in attrs.items(): e.set(qn(k.replace("_", ":", 1)), str(v))
    return e

def z(text):
    return insert_zwsp(text) if re.search(r"[฀-๿]", text) else text

def mk_run(text, bold=False, italic=False, vert=None, color=None, sz=None, font=None, zw=True):
    r = OxmlElement("w:r"); rpr = OxmlElement("w:rPr")
    if font:
        rf = OxmlElement("w:rFonts")
        for a in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"): rf.set(qn(a), font)
        rpr.append(rf)
    if bold: rpr.append(OxmlElement("w:b")); rpr.append(OxmlElement("w:bCs"))
    if italic: rpr.append(OxmlElement("w:i")); rpr.append(OxmlElement("w:iCs"))
    if color: c = OxmlElement("w:color"); c.set(qn("w:val"), color); rpr.append(c)
    if sz:
        a = OxmlElement("w:sz"); a.set(qn("w:val"), str(sz)); rpr.append(a); b = OxmlElement("w:szCs"); b.set(qn("w:val"), str(sz)); rpr.append(b)
    if vert: v = OxmlElement("w:vertAlign"); v.set(qn("w:val"), vert); rpr.append(v)
    lang = OxmlElement("w:lang"); lang.set(qn("w:bidi"), "th-TH"); rpr.append(lang)
    r.append(rpr)
    t = OxmlElement("w:t"); t.set("{http://www.w3.org/XML/1998/namespace}space", "preserve"); t.text = z(text) if zw else text
    r.append(t)
    return r

TAG = re.compile(r"<(b|i|sub|sup|ph|red)>(.*?)</\1>|<PH/>|<br/>", re.S)
def parse_inline(text, fmt=None, sz=None):
    fmt = dict(fmt or {}); out = []; pos = 0
    for m in TAG.finditer(text):
        if m.start() > pos: out.append(mk_run(text[pos:m.start()], sz=sz, **fmt))
        if m.group(0) == "<PH/>":
            out.append(mk_run("[Place Holder]", bold=True, color=RED, sz=sz, **{k: v for k, v in fmt.items() if k not in ("bold", "color")}))
        elif m.group(0) == "<br/>":
            r = OxmlElement("w:r"); r.append(OxmlElement("w:br")); out.append(r)
        else:
            tag, inner = m.group(1), m.group(2); f2 = dict(fmt)
            if tag == "b": f2["bold"] = True
            elif tag == "i": f2["italic"] = True
            elif tag == "sub": f2["vert"] = "subscript"
            elif tag == "sup": f2["vert"] = "superscript"
            elif tag in ("ph", "red"): f2["bold"] = True; f2["color"] = RED
            out.extend(parse_inline(inner, f2, sz))
        pos = m.end()
    if pos < len(text): out.append(mk_run(text[pos:], sz=sz, **fmt))
    return out

def mk_p(text="", style=None, first=False, after=None, before=None, jc=None, keep_next=False, left=None, hanging=None, page_break=False, tabs=None, sz=None, fmt=None, keep_lines=False):
    p = OxmlElement("w:p"); ppr = OxmlElement("w:pPr")
    if style: ps = OxmlElement("w:pStyle"); ps.set(qn("w:val"), style); ppr.append(ps)
    if keep_next: ppr.append(OxmlElement("w:keepNext"))
    if keep_lines: ppr.append(OxmlElement("w:keepLines"))
    if page_break: ppr.append(OxmlElement("w:pageBreakBefore"))
    if tabs:
        tb = OxmlElement("w:tabs")
        for kind, pos in tabs:
            t = OxmlElement("w:tab"); t.set(qn("w:val"), kind); t.set(qn("w:pos"), str(pos)); tb.append(t)
        ppr.append(tb)
    if after is not None or before is not None:
        sp = OxmlElement("w:spacing")
        if before is not None: sp.set(qn("w:before"), str(before))
        if after is not None: sp.set(qn("w:after"), str(after))
        ppr.append(sp)
    if first or left is not None:
        ind = OxmlElement("w:ind")
        if left is not None: ind.set(qn("w:left"), str(left))
        if hanging is not None: ind.set(qn("w:hanging"), str(hanging))
        if first: ind.set(qn("w:firstLine"), "720")
        ppr.append(ind)
    if jc: j = OxmlElement("w:jc"); j.set(qn("w:val"), jc); ppr.append(j)
    if len(ppr): p.append(ppr)
    for r in parse_inline(text, fmt, sz): p.append(r)
    return p

CITE = re.compile(r"\{c:([^}]+)\}")
REF = re.compile(r"\{(fig|tab|eq|fign|tabn|eqn):(\w+)\}")
def compress(nums):
    nums = sorted(set(nums)); out = []; i = 0
    while i < len(nums):
        j = i
        while j + 1 < len(nums) and nums[j + 1] == nums[j] + 1: j += 1
        if j - i >= 2: out.append(f"{nums[i]}\u2013{nums[j]}")
        else: out.extend(str(n) for n in nums[i:j + 1])
        i = j + 1
    return ", ".join(out)

class Registry:
    def __init__(self):
        self.cites = []; self.labels = {"fig": {}, "tab": {}, "eq": {}}; self.counters = {}
        self.chapter = "0"; self.missing = set()
    def cite_no(self, key):
        if key not in self.cites: self.cites.append(key)
        return self.cites.index(key) + 1

class Builder:
    def __init__(self, src, reg=None, dry=False):
        self.dry = dry; self.reg = reg or Registry(); self.final = reg is not None and not dry
        self.doc = Document(src); self.body = self.doc.element.body
        self.sect = self.body.find(qn("w:sectPr"))
        self.orig = [copy.deepcopy(ch) for ch in self.body.iterchildren() if ch.tag != qn("w:sectPr")]
        for ch in list(self.body.iterchildren()):
            if ch.tag != qn("w:sectPr"): self.body.remove(ch)
        self.toc_items = []        # (kind, level, text, bookmark)
    # ---- basic appenders ----
    def add(self, el):
        if self.dry: return el
        self.sect.addprevious(el); return el
    def keep(self, i, n=1):
        if self.dry: return
        for k in range(i, i + n): self.add(copy.deepcopy(self.orig[k]))
    def set_chapter(self, c): self.reg.chapter = str(c)
    def _label(self, kind, key):
        n = self.reg.counters.get((kind, self.reg.chapter), 0) + 1; self.reg.counters[(kind, self.reg.chapter)] = n
        lab = f"{self.reg.chapter}-{n}"; self.reg.labels[kind][key] = lab; return lab
    def R(self, t):
        t = str(t)
        def c(m):
            nums = []
            for k in m.group(1).split(","):
                k = k.strip()
                nums.append(self.reg.cite_no(k))
            return "[" + compress(nums) + "]"
        t = CITE.sub(c, t)
        def r(m):
            kind, key = m.groups(); bare = kind.endswith("n") and kind != "eq"; bare = bare or kind == "eqn"; kind = kind[:-1] if bare else kind
            lab = self.reg.labels[kind].get(key)
            if lab is None:
                if self.final: self.reg.missing.add(f"{kind}:{key}")
                return "?"
            if os.environ.get("REFDEBUG"): lab = "\u27e6" + lab + "\u27e7"
            return ("" if bare else {"fig": "ภาพที่ ", "tab": "ตารางที่ ", "eq": "สมการที่ "}[kind]) + lab
        return REF.sub(r, t)
    def p(self, text, **kw):
        kw.setdefault("after", 120); kw.setdefault("first", True)
        return self.add(mk_p(self.R(text), **kw))
    def plain(self, text, **kw): return self.add(mk_p(self.R(text), **kw))
    def bullets(self, items, ind=907, hang=340, after=40):
        for t in items: self.add(mk_p("•  " + self.R(t), left=ind, hanging=hang, after=after))
        self.add(mk_p("", after=60))
    def numbered(self, items, start=1):
        for k, t in enumerate(items, start): self.add(mk_p(f"{k}.  " + self.R(t), left=907, hanging=340, after=40))
        self.add(mk_p("", after=60))
    def bm(self, p, name):
        _bm[0] += 1
        s = OxmlElement("w:bookmarkStart"); s.set(qn("w:id"), str(_bm[0])); s.set(qn("w:name"), name)
        e = OxmlElement("w:bookmarkEnd"); e.set(qn("w:id"), str(_bm[0]))
        ppr = p.find(qn("w:pPr")); idx = 0 if ppr is None else list(p).index(ppr) + 1
        p.insert(idx, s); p.append(e)
    def h1(self, no, title, key=None):
        self.set_chapter(no)
        p = OxmlElement("w:p"); ppr = OxmlElement("w:pPr")
        ps = OxmlElement("w:pStyle"); ps.set(qn("w:val"), "Heading1"); ppr.append(ps); ppr.append(OxmlElement("w:pageBreakBefore"))
        sp = OxmlElement("w:spacing"); sp.set(qn("w:before"), "0"); ppr.append(sp); p.append(ppr)
        p.append(mk_run(f"บทที่ {no}", bold=True, font=HF)); r = OxmlElement("w:r"); r.append(OxmlElement("w:br")); p.append(r); p.append(mk_run(title, bold=True, font=HF))
        name = f"_TocPRP{len(self.toc_items)+1}"; self.bm(p, name); self.toc_items.append(("h1", 1, f"บทที่ {no} {title}", name)); return self.add(p)
    def h1_plain(self, title, toc=True):
        p = OxmlElement("w:p"); ppr = OxmlElement("w:pPr"); ps = OxmlElement("w:pStyle"); ps.set(qn("w:val"), "Heading1"); ppr.append(ps); ppr.append(OxmlElement("w:pageBreakBefore"))
        sp = OxmlElement("w:spacing"); sp.set(qn("w:before"), "0"); ppr.append(sp); p.append(ppr); p.append(mk_run(title, bold=True, font=HF))
        name = f"_TocPRP{len(self.toc_items)+1}"; self.bm(p, name)
        if toc: self.toc_items.append(("h1", 1, title, name))
        return self.add(p)
    def h2(self, text):
        p = mk_p(self.R(text), style="Heading2", fmt={"bold": True, "font": HF})
        name = f"_TocPRP{len(self.toc_items)+1}"; self.bm(p, name); self.toc_items.append(("h2", 2, self.R(text), name)); return self.add(p)
    def h3(self, text): return self.add(mk_p(self.R(text), before=160, after=80, first=True, keep_next=True, fmt={"bold": True}))
    # ---- figures ----
    def fig(self, key, path, width_cm, caption, source=None, note=None):
        lab = self._label("fig", key); caption = f"ภาพที่ {lab} " + caption
        if self.dry:
            self.R(caption); self.R(source or ""); self.R(note or ""); return
        p = self.doc.add_paragraph(); self.sect.addprevious(p._p)
        ppr = p._p.get_or_add_pPr()
        for tag, attrs in (("w:keepNext", {}), ("w:spacing", {"w:before": "160", "w:after": "40"}), ("w:jc", {"w:val": "center"})):
            e = OxmlElement(tag)
            for k, v in attrs.items(): e.set(qn(k), v)
            ppr.append(e)
        p.add_run().add_picture(path, width=Cm(width_cm))
        cp = mk_p(self.R(caption), style="FigCaption"); name = f"_TocPRP{len(self.toc_items)+1}"; self.bm(cp, name)
        self.toc_items.append(("fig", 1, self.R(caption), name)); self.add(cp)
        if source: self.add(mk_p(self.R(source), after=60, jc="center", sz=26))
        if note: self.add(mk_p(self.R(note), after=160, jc="center", sz=26))
    # ---- tables ----
    def table(self, key, caption, header, rows, widths, align=None, sz=22, shade_cols=None, bold_first_col=False, note=None, header_rows=1, row_shade=None, keep_all=False):
        lab = self._label("tab", key); caption = f"ตารางที่ {lab} " + caption
        if self.dry:
            self.R(caption); self.R(note or "")
            for h in (header if header and isinstance(header[0], (list, tuple)) else [header]):
                for x in h: self.R(x[0] if isinstance(x, tuple) else x)
            for row in rows:
                for x in row: self.R(x[1] if isinstance(x, tuple) else x)
            return
        cp = mk_p(self.R(caption), style="TabCaption"); name = f"_TocPRP{len(self.toc_items)+1}"; self.bm(cp, name)
        self.toc_items.append(("tab", 1, self.R(caption), name)); self.add(cp)
        tbl = OxmlElement("w:tbl"); pr = OxmlElement("w:tblPr")
        ts = OxmlElement("w:tblStyle"); ts.set(qn("w:val"), "TableGrid"); pr.append(ts)
        tw = OxmlElement("w:tblW"); tw.set(qn("w:w"), str(sum(widths))); tw.set(qn("w:type"), "dxa"); pr.append(tw)
        j = OxmlElement("w:jc"); j.set(qn("w:val"), "center"); pr.append(j)
        lay = OxmlElement("w:tblLayout"); lay.set(qn("w:type"), "fixed"); pr.append(lay)
        look = OxmlElement("w:tblLook"); look.set(qn("w:val"), "04A0"); pr.append(look); tbl.append(pr)
        grid = OxmlElement("w:tblGrid")
        for w_ in widths: g = OxmlElement("w:gridCol"); g.set(qn("w:w"), str(w_)); grid.append(g)
        tbl.append(grid)
        align = align or ["center"] * len(widths)
        def cell(txt, w_, fill=None, bold=False, jc="center", span=1):
            tc = OxmlElement("w:tc"); tcp = OxmlElement("w:tcPr")
            cw = OxmlElement("w:tcW"); cw.set(qn("w:w"), str(w_)); cw.set(qn("w:type"), "dxa"); tcp.append(cw)
            if span > 1: gs = OxmlElement("w:gridSpan"); gs.set(qn("w:val"), str(span)); tcp.append(gs)
            if fill: sh = OxmlElement("w:shd"); sh.set(qn("w:val"), "clear"); sh.set(qn("w:color"), "auto"); sh.set(qn("w:fill"), fill); tcp.append(sh)
            va = OxmlElement("w:vAlign"); va.set(qn("w:val"), "center"); tcp.append(va); tc.append(tcp)
            lines = str(txt).split("\n")
            for ln in lines:
                tc.append(mk_p(self.R(ln), after=20, before=20, jc=jc, sz=sz, fmt={"bold": True} if bold else None))
            return tc
        hdrs = header if header and isinstance(header[0], (list, tuple)) else [header]
        for hrow in hdrs:
            tr = OxmlElement("w:tr"); trp = OxmlElement("w:trPr"); trp.append(OxmlElement("w:cantSplit")); trp.append(OxmlElement("w:tblHeader")); tr.append(trp)
            ci = 0
            for h in hrow:
                if isinstance(h, tuple): txt, span = h
                else: txt, span = h, 1
                tr.append(cell(txt, sum(widths[ci:ci + span]), fill="EDEDED", bold=True, span=span)); ci += span
            tbl.append(tr)
        for ri, row in enumerate(rows):
            tr = OxmlElement("w:tr"); trp = OxmlElement("w:trPr"); trp.append(OxmlElement("w:cantSplit")); tr.append(trp)
            if row and isinstance(row[0], tuple) and row[0][0] == "__section__":      # full-width section row
                tr.append(cell(row[0][1], sum(widths), fill="F5F5F5", bold=True, jc="left", span=len(widths))); tbl.append(tr); continue
            for ci, v in enumerate(row):
                fill = ("F7F7F7" if shade_cols and ci in shade_cols else None)
                if row_shade and ri in row_shade: fill = row_shade[ri]
                tr.append(cell(v, widths[ci], fill=fill, bold=(bold_first_col and ci == 0), jc=align[ci]))
            tbl.append(tr)
        self.add(tbl)
        self.add(mk_p(self.R(note) if note else "", after=100, sz=26 if note else None, left=0))
    def tabimg(self, key, caption, path, width_cm, note=None):
        lab = self._label("tab", key); caption = f"ตารางที่ {lab} " + caption
        if self.dry: self.R(caption); return
        cp = mk_p(self.R(caption), style="TabCaption"); name = f"_TocPRP{len(self.toc_items)+1}"; self.bm(cp, name)
        self.toc_items.append(("tab", 1, self.R(caption), name)); self.add(cp)
        p = self.doc.add_paragraph(); self.sect.addprevious(p._p); ppr = p._p.get_or_add_pPr()
        for tag, attrs in (("w:spacing", {"w:before": "40", "w:after": "40"}), ("w:jc", {"w:val": "center"})):
            e = OxmlElement(tag)
            for k, v in attrs.items(): e.set(qn(k), v)
            ppr.append(e)
        p.add_run().add_picture(path, width=Cm(width_cm))
        if note: self.add(mk_p(self.R(note), after=100, jc="center", sz=26))
    # ---- equations ----
    def eq(self, key, latex):
        label = self._label("eq", key)
        if self.dry: return
        p = OxmlElement("w:p"); ppr = OxmlElement("w:pPr")
        tabs = OxmlElement("w:tabs")
        for kind, pos in (("center", 4300), ("right", 8600)):
            t = OxmlElement("w:tab"); t.set(qn("w:val"), kind); t.set(qn("w:pos"), str(pos)); tabs.append(t)
        ppr.append(tabs); sp = OxmlElement("w:spacing"); sp.set(qn("w:before"), "120"); sp.set(qn("w:after"), "120"); ppr.append(sp); p.append(ppr)
        r = OxmlElement("w:r"); r.append(OxmlElement("w:tab")); p.append(r)
        om = latex_to_omml(latex)
        if om.tag.endswith("oMathPara"): om = om[0]
        p.append(om)
        r2 = OxmlElement("w:r"); r2.append(OxmlElement("w:tab")); p.append(r2)
        p.append(mk_run(f"สมการที่ {label}"))
        return self.add(p)
    def pagebreak(self):
        if self.dry: return
        p = OxmlElement("w:p"); r = OxmlElement("w:r"); b = OxmlElement("w:br"); b.set(qn("w:type"), "page"); r.append(b); p.append(r); return self.add(p)
    # ---- finalize ----
    def prune_rels(self):
        xml = etree.tostring(self.body).decode()
        used = set(re.findall(r'r:(?:embed|id|link)="(rId\d+)"', xml))
        for sect_el in self.body.iter(qn("w:sectPr")):
            for ref in list(sect_el):
                rid = ref.get("{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id")
                if rid: used.add(rid)
        part = self.doc.part; dropped = 0
        for rid, rel in list(part.rels.items()):
            if rel.reltype.endswith("/image") and rid not in used:
                part.rels.pop(rid); dropped += 1
        return dropped
    def set_update_fields(self):
        st = self.doc.settings.element
        if st.find(qn("w:updateFields")) is None:
            u = OxmlElement("w:updateFields"); u.set(qn("w:val"), "true"); st.append(u)
    def save(self, path):
        self.prune_rels(); self.doc.save(path)


# ---------------------------------------------------------------- front-matter helpers (TOC / lists)
def _fld(kind):
    r = OxmlElement("w:r"); f = OxmlElement("w:fldChar"); f.set(qn("w:fldCharType"), kind); r.append(f); return r

def _instr(text):
    r = OxmlElement("w:r"); i = OxmlElement("w:instrText"); i.set("{http://www.w3.org/XML/1998/namespace}space", "preserve"); i.text = text; r.append(i); return r

def toc_paragraphs(instr, entries, pages, indent2=360):
    """entries = [(kind, level, text, bookmark)]; returns list of <w:p> forming one TOC-type field with cached entries."""
    out = []
    for n, (kind, level, text, bm) in enumerate(entries):
        p = OxmlElement("w:p"); ppr = OxmlElement("w:pPr")
        ps = OxmlElement("w:pStyle"); ps.set(qn("w:val"), "TOC1"); ppr.append(ps)
        tabs = OxmlElement("w:tabs"); t = OxmlElement("w:tab"); t.set(qn("w:val"), "right"); t.set(qn("w:leader"), "dot"); t.set(qn("w:pos"), "8659"); tabs.append(t); ppr.append(tabs)
        sp = OxmlElement("w:spacing"); sp.set(qn("w:after"), "40" if kind in ("fig", "tab") else "60"); ppr.append(sp)
        if level == 2 and kind == "h2":
            ind = OxmlElement("w:ind"); ind.set(qn("w:left"), str(indent2)); ppr.append(ind)
        elif kind in ("fig", "tab"):
            ind = OxmlElement("w:ind"); ind.set(qn("w:left"), "1134"); ind.set(qn("w:hanging"), "1134"); ppr.append(ind)
        p.append(ppr)
        if n == 0:
            p.append(_fld("begin")); p.append(_instr(instr)); p.append(_fld("separate"))
        h = OxmlElement("w:hyperlink"); h.set(qn("w:anchor"), bm); h.set(qn("w:history"), "1")
        for r in parse_inline(text, {"bold": kind == "h1"}, sz=(28 if kind in ("fig", "tab") else None)):
            rpr = r.find(qn("w:rPr")); rs = OxmlElement("w:rStyle"); rs.set(qn("w:val"), "Hyperlink"); rpr.insert(0, rs)
            np_ = OxmlElement("w:noProof"); nxt = next((c for c in rpr if c.tag in {qn("w:" + t) for t in ("color", "spacing", "sz", "szCs", "vertAlign", "lang")}), None)
            (nxt.addprevious(np_) if nxt is not None else rpr.append(np_)); h.append(r)
        r2 = OxmlElement("w:r"); r2p = OxmlElement("w:rPr"); r2p.append(OxmlElement("w:noProof")); r2p.append(OxmlElement("w:webHidden")); r2.append(r2p); r2.append(OxmlElement("w:tab")); h.append(r2)
        r3 = OxmlElement("w:r"); r3p = OxmlElement("w:rPr"); r3p.append(OxmlElement("w:noProof")); r3p.append(OxmlElement("w:webHidden")); r3.append(r3p)
        tt = OxmlElement("w:t"); tt.text = str(pages.get(text, "")); r3.append(tt); h.append(r3)
        p.append(h); out.append(p)
    endp = OxmlElement("w:p"); endp.append(_fld("end")); out.append(endp)
    return out
