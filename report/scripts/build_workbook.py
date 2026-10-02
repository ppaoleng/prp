"""Analysis workbook for the PRP 2569 report (no .xlsx was supplied with the draft, so this one is built from the draft's tables).
Summary sheets use live Excel formulas; placeholder (non-measured) values are written in bold red and marked [Place Holder]."""
import os, sys
HERE = os.path.dirname(os.path.abspath(__file__)); sys.path.insert(0, HERE)
import pandas as pd, numpy as np
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.chart import BarChart, LineChart, Reference, Series
from openpyxl.utils import get_column_letter as L
import rnum
from prp_data import *
import lca_model

OUT = os.path.join(HERE, "..", "output", "PRP2569_analysis.xlsx")
N = rnum.load()
RED = Font(bold=True, color="C00000"); BOLD = Font(bold=True); HDR = PatternFill("solid", fgColor="E8E8E4"); IN = PatternFill("solid", fgColor="FFF8DC")
thin = Side(style="thin", color="BBBBBB"); BOX = Border(left=thin, right=thin, top=thin, bottom=thin)

def header(ws, row, cols):
    for j, c in enumerate(cols, 1):
        x = ws.cell(row=row, column=j, value=c); x.font = BOLD; x.fill = HDR; x.border = BOX; x.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")

def widths(ws, ws_w):
    for j, w in enumerate(ws_w, 1): ws.column_dimensions[L(j)].width = w

wb = Workbook()
# ---------------------------------------------------------------- README
ws = wb.active; ws.title = "README"
rows = [
    ("PRP 2569 — analysis workbook (v2 of the full report)", None),
    ("Source", "Built from the tables of the draft '2.PRP_รายงานฉบับสมบูรณ์ 2569 [DRAFT].docx'. No .xlsx was supplied with the draft."),
    ("Status colours", "Black = value recorded in the draft.  Bold red = [Place Holder] (not measured / assumed; must be replaced before publication).  Yellow cells = editable inputs (emission factors, price ratio r)."),
    ("Strength_raw", "Per-specimen compressive strength (MPa). Three draft cells have a number but no original record: they are NOT in the specimen columns (so AVERAGE/STDEV ignore them); their draft values are listed in the red 'PH' columns."),
    ("Strength_summary", "Live formulas: n, mean, SD, CV, difference to the reported mean, and 28-d % change vs control."),
    ("Density / Absorption", "Per-specimen density = mass / volume and absorption = (wet-dry)/dry x 100 (formulas)."),
    ("Maturity", "EN 1992-1-1 maturity factor beta(t)=exp{s[1-(28/t)^0.5]}: s is a yellow input; compare predicted with measured f(t)/f(28)."),
    ("LCA", "Screening cradle-to-gate GWP (functional unit 1 m3 of block).  Emission factors are ASSUMED (yellow, red); change them and every result updates."),
    ("Cost", "Relative material cost vs RA as a function of the biochar price ratio r (yellow input)."),
    ("BPN_PLACEHOLDER", "Hypothetical British Pendulum Numbers (all red). NOT measurements."),
    ("Charts", "Native Excel charts built from the sheets above."),
    ("Reproducibility", "All numbers are produced by the scripts in report/scripts (seeds fixed). See report/docs/CHANGELOG_and_review_notes.md."),
]
for i, (a, b) in enumerate(rows, 1):
    ws.cell(row=i, column=1, value=a).font = BOLD
    if b: ws.cell(row=i, column=2, value=b).alignment = Alignment(wrap_text=True, vertical="top")
widths(ws, [24, 120])

# ---------------------------------------------------------------- Strength_raw
L_ = strength_long()
ws = wb.create_sheet("Strength_raw")
header(ws, 1, ["series", "mix", "age_d", "spec1_MPa", "spec2_MPa", "spec3_MPa", "reported_mean", "PH spec1", "PH spec2", "PH spec3"])
r = 2; idx = {}
for (series, mix, age), g in L_.groupby(["series", "mix", "age"], sort=False):
    ws.cell(row=r, column=1, value="I" if series == "I" else "II"); ws.cell(row=r, column=2, value=mix); ws.cell(row=r, column=3, value=int(age))
    for k, spec in enumerate(("s1", "s2", "s3")):
        x = g[g.spec == spec].iloc[0]
        if x.placeholder:
            c = ws.cell(row=r, column=8 + k, value=float(x.fc)); c.font = RED
        else:
            ws.cell(row=r, column=4 + k, value=float(x.fc))
    ws.cell(row=r, column=7, value=float(g.mean_reported.iloc[0]))
    idx[(mix, age)] = r; r += 1
widths(ws, [8, 10, 8, 12, 12, 12, 14, 10, 10, 10]); ws.freeze_panes = "A2"

# ---------------------------------------------------------------- Strength_summary
ss = wb.create_sheet("Strength_summary")
header(ss, 1, ["series", "mix", "age_d", "n", "mean (recomputed)", "SD", "CV %", "reported mean", "recomputed - reported", "28-d % vs series control"])
for (mix, age), rr in idx.items():
    row = rr
    ss.cell(row=row, column=1, value=f"=Strength_raw!A{rr}"); ss.cell(row=row, column=2, value=f"=Strength_raw!B{rr}"); ss.cell(row=row, column=3, value=f"=Strength_raw!C{rr}")
    ss.cell(row=row, column=4, value=f"=COUNT(Strength_raw!D{rr}:F{rr})")
    ss.cell(row=row, column=5, value=f"=AVERAGE(Strength_raw!D{rr}:F{rr})")
    ss.cell(row=row, column=6, value=f"=STDEV(Strength_raw!D{rr}:F{rr})")
    ss.cell(row=row, column=7, value=f"=F{row}/E{row}*100")
    ss.cell(row=row, column=8, value=f"=Strength_raw!G{rr}")
    ss.cell(row=row, column=9, value=f"=E{row}-H{row}")
    ctrl = idx[("RA", 28)] if mix.startswith(("Rab", "RA")) or mix == "NA" else idx[("RCA0", 28)]
    if age == 28: ss.cell(row=row, column=10, value=f"=(H{row}/Strength_raw!G{ctrl}-1)*100")
    for c in (5, 6, 7, 9, 10): ss.cell(row=row, column=c).number_format = "0.00"
widths(ss, [8, 10, 8, 6, 18, 8, 8, 14, 20, 22]); ss.freeze_panes = "A2"
# compact 28-d table for charts
ss.cell(row=len(idx) + 4, column=1, value="28-d reported mean (for chart)").font = BOLD
c0 = len(idx) + 5
for i, m in enumerate(["NA", "RA", "Rab5", "Rab10", "Rab15", "RCA0", "RCA5", "RCA10", "RCA15"]):
    ss.cell(row=c0 + i, column=2, value=m); ss.cell(row=c0 + i, column=3, value=f"=Strength_raw!G{idx[(m, 28)]}")
ch = BarChart(); ch.type = "col"; ch.title = "28-d compressive strength (reported means, MPa)"; ch.y_axis.title = "MPa"
ch.add_data(Reference(ss, min_col=3, min_row=c0, max_row=c0 + 8), titles_from_data=False); ch.set_categories(Reference(ss, min_col=2, min_row=c0, max_row=c0 + 8)); ch.legend = None; ch.height = 8; ch.width = 16
ss.add_chart(ch, f"L2")

# ---------------------------------------------------------------- Density
D = density_long(); wd = wb.create_sheet("Density")
header(wd, 1, ["mix", "age_d", "mass_kg", "volume_m3", "density (formula, kg/m3)", "density printed in draft", "reported mean"])
for i, rr in enumerate(D.itertuples(), 2):
    wd.cell(row=i, column=1, value=rr.mix); wd.cell(row=i, column=2, value=int(rr.age_d)); wd.cell(row=i, column=3, value=float(rr.mass_kg)); wd.cell(row=i, column=4, value=float(rr.volume_m3))
    wd.cell(row=i, column=5, value=f"=C{i}/D{i}").number_format = "0"; wd.cell(row=i, column=6, value=float(rr.density))
    if pd.notna(rr.mean_reported): wd.cell(row=i, column=7, value=float(rr.mean_reported))
widths(wd, [10, 8, 10, 12, 24, 22, 14]); wd.freeze_panes = "A2"

# ---------------------------------------------------------------- Absorption
wa = wb.create_sheet("Absorption")
header(wa, 1, ["series", "mix", "age_d", "specimen", "wet_kg", "dry_kg", "absorption % (formula)", "absorption % printed", "flag"])
A = absorption_long(); i = 2
flag = {("RCA10", 14), ("RCA15", 14), ("RCA15", 28)}
for rr in A.itertuples():
    wa.cell(row=i, column=1, value=rr.series); wa.cell(row=i, column=2, value=rr.mix); wa.cell(row=i, column=3, value=int(rr.age)); wa.cell(row=i, column=4, value=str(rr.spec))
    if pd.notna(rr.wet):
        wa.cell(row=i, column=5, value=float(rr.wet)); wa.cell(row=i, column=6, value=float(rr.dry)); wa.cell(row=i, column=7, value=f"=(E{i}-F{i})/F{i}*100").number_format = "0.00"
    wa.cell(row=i, column=8, value=float(rr.ab))
    if (rr.mix, int(rr.age)) in flag:
        wa.cell(row=i, column=9, value="re-test needed (see Appendix B)").font = RED
    i += 1
widths(wa, [8, 10, 8, 10, 10, 10, 22, 20, 30]); wa.freeze_panes = "A2"

# ---------------------------------------------------------------- Maturity
wm = wb.create_sheet("Maturity")
header(wm, 1, ["mix", "f7/f28 measured", "f14/f28 measured", "s (input)", "beta(7) model", "beta(14) model", "7-d residual", "14-d residual"])
fit = pd.read_csv(os.path.join(DATA, "stats_maturity_fits.csv"), keep_default_na=False)
for i, rr in enumerate(fit.itertuples(), 2):
    wm.cell(row=i, column=1, value=rr.mix); wm.cell(row=i, column=2, value=float(rr.f7_f28)); wm.cell(row=i, column=3, value=float(rr.f14_f28))
    c = wm.cell(row=i, column=4, value=round(float(rr.s_fit), 3)); c.fill = IN
    wm.cell(row=i, column=5, value=f"=EXP(D{i}*(1-(28/7)^0.5))").number_format = "0.000"; wm.cell(row=i, column=6, value=f"=EXP(D{i}*(1-(28/14)^0.5))").number_format = "0.000"
    wm.cell(row=i, column=7, value=f"=B{i}-E{i}").number_format = "0.000"; wm.cell(row=i, column=8, value=f"=C{i}-F{i}").number_format = "0.000"
wm.cell(row=len(fit) + 3, column=1, value="EN 1992-1-1 reference s: 0.20 (rapid), 0.25 (normal), 0.38 (slow) — values as quoted in report chapter 2; confirm against the standard.").font = RED
widths(wm, [10, 16, 16, 10, 14, 14, 14, 14])

# ---------------------------------------------------------------- LCA
wl = wb.create_sheet("LCA")
wl.cell(row=1, column=1, value="Block volume (m3)").font = BOLD; wl.cell(row=1, column=2, value=BLOCK_VOLUME_M3)
wl.cell(row=2, column=1, value="Emission factors (kg CO2e per kg) — ASSUMED [Place Holder]").font = RED
header(wl, 3, ["material", "EF (mode)", "low", "high", "status"])
fac = lca_model.FACTORS; ef_row = {}
for i, k in enumerate(["cement", "sand", "stone", "rca", "biochar"], 4):
    wl.cell(row=i, column=1, value=k)
    for j, v in enumerate(fac[k]): c = wl.cell(row=i, column=2 + j, value=v); c.fill = IN; c.font = RED
    wl.cell(row=i, column=5, value="ASSUMED — replace with Thai LCI / TGO / ecoinvent value" if k != "cement" else "mode: Thai OPC 0.76 (Khongprom & Suwanmanee 2017); high = assumed")
    ef_row[k] = i
wl.cell(row=9, column=1, value="Carbon fraction of biochar"); c = wl.cell(row=9, column=2, value=lca_model.C_FRAC[1]); c.fill = IN; c.font = RED
wl.cell(row=10, column=1, value="Permanence over 100 y"); c = wl.cell(row=10, column=2, value=lca_model.PERMANENCE[1]); c.fill = IN; c.font = RED
wl.cell(row=11, column=1, value="44/12"); wl.cell(row=11, column=2, value="=44/12")
H = 13
header(wl, H, ["mix", "cement g/block", "sand g", "stone g", "RCA g", "biochar g", "cement kg/m3", "sand kg/m3", "stone kg/m3", "RCA kg/m3", "biochar kg/m3", "GWP A", "biochar credit", "GWP B", "cement share of A %", "f_c 28 d", "ci A", "ci B", "bi (kg cement / m3 / MPa)"])
batch = {"NA": (600, 1800, 900, 0, 0), "RA": (600, 900, 900, 900, 0), "Rab5": (570, 900, 900, 900, 45), "Rab10": (540, 900, 900, 900, 90), "Rab15": (510, 900, 900, 900, 135),
         "Rab20": (480, 900, 900, 900, 180), "Rab25": (450, 900, 900, 900, 225), "RCA0": (600, 900, 1800, 0, 48), "RCA5": (600, 855, 1800, 45, 48), "RCA10": (600, 810, 1800, 90, 48), "RCA15": (600, 765, 1800, 135, 48)}
fc = {m: N[f"fc28_{m}"] for m in ["NA", "RA", "Rab5", "Rab10", "Rab15", "RCA0", "RCA5", "RCA10", "RCA15"]}
lrow = {}
for i, (m, v) in enumerate(batch.items(), H + 1):
    wl.cell(row=i, column=1, value=m + ("*" if m in ("Rab20", "Rab25") else ""))
    for j, x in enumerate(v): wl.cell(row=i, column=2 + j, value=x)
    for j in range(5): wl.cell(row=i, column=7 + j, value=f"={L(2 + j)}{i}/1000/$B$1").number_format = "0.0"
    wl.cell(row=i, column=12, value=f"=G{i}*$B$4+H{i}*$B$5+I{i}*$B$6+J{i}*$B$7+K{i}*$B$8").number_format = "0.0"
    wl.cell(row=i, column=13, value=f"=K{i}*$B$9*$B$10*$B$11").number_format = "0.0"
    wl.cell(row=i, column=14, value=f"=L{i}-M{i}").number_format = "0.0"
    wl.cell(row=i, column=15, value=f"=G{i}*$B$4/L{i}*100").number_format = "0.0"
    if m in fc:
        wl.cell(row=i, column=16, value=fc[m]); wl.cell(row=i, column=17, value=f"=L{i}/P{i}").number_format = "0.00"; wl.cell(row=i, column=18, value=f"=N{i}/P{i}").number_format = "0.00"; wl.cell(row=i, column=19, value=f"=G{i}/P{i}").number_format = "0.00"
    lrow[m] = i
widths(wl, [14] + [13] * 18)
wl.cell(row=H + 14, column=1, value="* designed but not yet tested (no f_c). NA batch (600:1800:900) is inferred from the 1:3:1.5 ratio and is NOT printed in the draft [Place Holder].").font = RED
bc = BarChart(); bc.type = "col"; bc.title = "GWP, scenario A vs B (kg CO2e / m3) — screening"; bc.height = 8; bc.width = 18
for col, nm in ((12, "A (no credit)"), (14, "B (storage credit)")):
    s = Series(Reference(wl, min_col=col, min_row=H + 1, max_row=H + 11), title=nm); bc.series.append(s)
bc.set_categories(Reference(wl, min_col=1, min_row=H + 1, max_row=H + 11)); wl.add_chart(bc, f"A{H + 17}")

# ---------------------------------------------------------------- Cost
wc = wb.create_sheet("Cost")
wc.cell(row=1, column=1, value="Biochar price ratio r (biochar / cement, per kg) [Place Holder]").font = RED; c = wc.cell(row=1, column=2, value=0.5); c.fill = IN; c.font = RED
wc.cell(row=2, column=1, value="Sand/stone price (cement = 1)"); wc.cell(row=2, column=2, value=0.1).fill = IN
wc.cell(row=3, column=1, value="RCA price"); wc.cell(row=3, column=2, value=0.0).fill = IN
header(wc, 5, ["mix", "relative material cost (per m3)", "cost vs RA", "cost per MPa vs RA"])
for i, m in enumerate(["RA", "Rab5", "Rab10", "Rab15"], 6):
    li = lrow[m]
    wc.cell(row=i, column=1, value=m)
    wc.cell(row=i, column=2, value=f"=LCA!G{li}*1+(LCA!H{li}+LCA!I{li})*$B$2+LCA!J{li}*$B$3+LCA!K{li}*$B$1").number_format = "0.0"
    wc.cell(row=i, column=3, value=f"=B{i}/$B$6").number_format = "0.000"
    wc.cell(row=i, column=4, value=f"=(B{i}/LCA!P{li})/($B$6/LCA!P{lrow['RA']})").number_format = "0.000"
wc.cell(row=11, column=1, value="Break-even r (Series I) = cement removed / biochar added = 30/45 = 0.667")
widths(wc, [52, 30, 14, 20])

# ---------------------------------------------------------------- BPN placeholder
B = N["bpn"]; wp = wb.create_sheet("BPN_PLACEHOLDER")
wp.cell(row=1, column=1, value="[Place Holder] — hypothetical BPN values; NOT measurements. 28-d anchors = the draft's example values (confirmed as examples by the researcher).").font = RED
header(wp, 2, ["series", "mix", "age_d", "BPN dry", "SD dry", "BPN wet", "SD wet", "dry - wet", "status"])
for i, rr in enumerate(B.itertuples(), 3):
    vals = [rr.series, rr.mix, int(rr.age), rr.bpn_dry, rr.sd_dry, rr.bpn_wet, rr.sd_wet]
    for j, v in enumerate(vals, 1):
        c = wp.cell(row=i, column=j, value=v if not isinstance(v, (np.floating, np.integer)) else float(v)); c.font = RED
    wp.cell(row=i, column=8, value=f"=D{i}-F{i}").font = RED; wp.cell(row=i, column=9, value="[Place Holder]").font = RED
widths(wp, [8, 10, 8, 10, 10, 10, 10, 10, 16])
lc = LineChart(); lc.title = "[Place Holder] wet BPN at 28 d"; lc.height = 7; lc.width = 14
widths(wp, [8, 10, 8, 10, 10, 10, 10, 10, 16])

os.makedirs(os.path.dirname(OUT), exist_ok=True)
wb.save(OUT); print("saved", OUT)
