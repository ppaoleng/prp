import json, docx
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
d = docx.Document()
for s in d.sections: s.left_margin = s.right_margin = Cm(1.9); s.top_margin = s.bottom_margin = Cm(1.8)
st = d.styles["Normal"]; st.font.name = "TH Sarabun New"; st.font.size = Pt(14)
st.element.rPr.rFonts.set(qn("w:eastAsia"), "TH Sarabun New"); st.element.rPr.rFonts.set(qn("w:cs"), "TH Sarabun New")
def H(t, lvl=1):
    p = d.add_heading(t, lvl)
    for r in p.runs: r.font.name = "TH Sarabun New"; r.font.color.rgb = RGBColor(0x1a, 0x1a, 0x1a); r._r.rPr.rFonts.set(qn("w:cs"), "TH Sarabun New")
def P(t, bold=False, red=False):
    p = d.add_paragraph(); r = p.add_run(t); r.bold = bold
    if red: r.font.color.rgb = RGBColor(255, 0, 0)
    return p
def T(header, rows, widths):
    t = d.add_table(rows=1, cols=len(header)); t.style = "Table Grid"
    for i, h in enumerate(header):
        c = t.rows[0].cells[i]; c.text = ""; r = c.paragraphs[0].add_run(h); r.bold = True; r.font.size = Pt(12)
    for row in rows:
        cells = t.add_row().cells
        for i, v in enumerate(row):
            cells[i].text = ""; r = cells[i].paragraphs[0].add_run(v); r.font.size = Pt(12)
    for row in t.rows:
        for i, w in enumerate(widths): row.cells[i].width = Cm(w)
    d.add_paragraph()
H("บันทึกการแก้ไขเปเปอร์ Bond_RB_GFRP (v2 → v3)", 0)
P("ไฟล์ต้นฉบับ v2 ไม่ได้ถูกแก้ไข (เก็บสำเนาไว้เป็น Bond_RB_GFRP_manuscript_v2_original.docx) ไฟล์ใหม่คือ Bond_RB_GFRP_manuscript_v3.docx เนื้อหาที่เพิ่มใหม่ทั้งหมดใช้ตัวอักษรสีแดง ค้นหาและลบออกได้ง่าย")
H("1. สิ่งที่ต้องรู้ก่อนใช้ไฟล์", 1)
for t in [
 "ตาราง S1–S4 ท้ายเล่ม (Appendix S) เป็นค่าสมมติทั้งหมด ตามที่ขอ เขียนให้อยู่ในช่วงที่สมเหตุสมผลสำหรับการทดสอบประเภทนี้ ไม่ใช่ข้อมูลที่วัดจริง ห้ามนำไปอ้างอิง คำนวณ หรือส่งวารสาร ต้องแทนด้วยผลทดสอบจริงหรือลบทิ้ง ตาราง S5 เป็นแบบฟอร์มเปล่า ไม่มีข้อมูลใด ๆ",
 "ตัวเลขในผลการทดลองหลัก (Section 3, Abstract, Conclusions, Table 1–10) ไม่ถูกเปลี่ยน ค่าสมมติไม่ได้ถูกนำไปใช้ในการวิเคราะห์ใด ๆ",
 "ไม่มีข้อมูลดิบรายตัวอย่างในไฟล์ที่แนบมา รูป Fig. 3, 4, 5, 9 จึงวาดใหม่จากค่าเฉลี่ยและ SD ของแต่ละกลุ่ม โดยค่าเฉลี่ยของ epoxy 9 กลุ่มใช้ตัวเลขในข้อความ Section 3.3 ส่วนค่าอื่นอ่านจากรูป Fig. 3 ของ v2 แล้วตรวจสอบย้อนกลับกับ Table 3: ค่าเฉลี่ยรวมของ epoxy และ grout ตรงกันภายใน 0.01 MPa และ SD รวมตรงกันภายใน 0.02 MPa ส่วน mortar ตรวจ SD ไม่ได้เพราะ SD ภายในกลุ่มซ่อนอยู่ใต้สัญลักษณ์ (ค่าเฉลี่ยตรงกับข้อความ Section 3.3) ข้อจำกัด: จุดข้อมูลรายตัวอย่าง (light points) หายไปจากรูปเหล่านี้ เพราะไม่มีข้อมูลดิบ และ SD ของ mortar ไม่ได้ขีดเป็น error bar (น้อยกว่า 0.1 MPa ซ่อนอยู่ในสัญลักษณ์) เมื่อได้ไฟล์ดิบแล้วควรวาดใหม่อีกครั้ง (ใส่ไว้ในหมายเหตุสีแดงที่ caption Fig. 3 แล้ว)",
 "เอกสารความเห็น reviewer ที่แนบมามีเพียง 4 ข้อ (ข้อ 1, 2, 4, 7) จากหัวข้อที่ระบุว่า 10 ข้อ ข้อ 3, 5, 6, 8, 9, 10 ยังไม่ได้ดำเนินการเพราะไม่มีเนื้อหา",
 "รูป Fig. 2(c) ยังเห็นขาและรองเท้าของผู้ปฏิบัติงาน และ (d)(e) เห็นมือสวมถุงมือ ต้องได้รับความยินยอมเป็นลายลักษณ์อักษรตามที่ v2 ระบุไว้แล้ว หรือครอปรูปเพิ่มก่อนส่ง ตัวหนังสือที่ฝังอยู่ในภาพถ่ายเดิมถูกตัดออก ย้ายไปอยู่ใน caption แทน",
]: d.add_paragraph(t, style="List Bullet")
H("2. การดำเนินการตามความเห็นของ reviewer", 1)
T(["ข้อ", "ประเด็น", "สิ่งที่ทำใน v3", "สถานะ"], [
 ["1", "ไม่ได้วัด f′c จริง", "ย่อหน้าสีแดงใน Sec. 2.2 ชี้ไปที่ Table S1 (ผลลูกบาศก์ข้างเคียง ค่าสมมติ) แก้ข้อความ 'nominal grade labels and not measured values' ใน Sec. 2.2 ปรับประโยคสุดท้ายของ Abstract และเพิ่มหมายเหตุสีแดงใน Sec. 4.8", "ต้องมีผลทดสอบจริง (หล่อและกดลูกบาศก์ชุดใหม่) ก่อนปรับข้อสรุปเรื่อง f′c"],
 ["2", "ที่มาของข้อมูล / raw data", "เพิ่ม Table S5 (แบบตรวจสอบที่มาของข้อมูล ว่างเปล่า) หมายเหตุสีแดงใน Sec. 2.1 และ Data availability แก้ 'confirmatory' เป็น 'inferential' ใน Sec. 2.1 เพราะขัดกับประโยคที่ว่าค่า p เป็นเพียงค่าเชิงพรรณนา", "ต้องกู้ไฟล์ดิบจากเครื่อง UTM / รูปถ่าย ถ้าไม่ได้ ให้ตัดกลุ่มที่ไม่มีหลักฐาน"],
 ["4", "วิธีทดสอบไม่เป็นมาตรฐานปัจจุบัน", "ย่อหน้าสีแดงใน Sec. 2.4 และ 4.6 ชี้ไปที่ Table S4 (ชุด l_e = 5 d_b ตาม ACI 440.3R Annex B.3 ค่าสมมติ) แก้ถ้อยคำ 'bond capacity' เป็น 'pull-out capacity' และ 'Slip records are not analysed here' ให้ชัดเจน เพิ่มแผนภาพ Fig. 1 กล่องเส้นประ", "ต้องทดสอบจริง และนำ load–slip จาก LVDT มาวิเคราะห์ถ้ามีไฟล์"],
 ["7", "ไม่ทราบสมบัติเหล็กและ GFRP", "Table S2 (เหล็ก SR24 coupon) และ Table S3 (GFRP) ค่าสมมติ ย่อหน้าสีแดงใน Sec. 2.2 แก้ Sec. 3.4 จาก 'approached' เป็น 'reached or exceeded' เพราะมี 8 จาก 45 ตัวอย่างเกิน 235 MPa", "ต้องทดสอบ tensile coupon จริง และ Sec. 3.4 / 4.3 ต้องเขียนใหม่ตามค่าจริง"],
], [1.0, 3.6, 8.0, 4.6])
H("3. รูปที่ออกแบบใหม่", 1)
T(["รูป", "การเปลี่ยนแปลง"], [
 ["Fig. 1", "Flowchart จัดใหม่เป็น 3 แถว (เหล็กเรียบ / GFRP / งานที่เพิ่มในการแก้ไข) 5 ขั้นตอน ตัวอักษรเดียวกันทั้งรูป กล่องเส้นประแสดงงานที่ reviewer ขอให้เพิ่ม (cube, coupon, l_e = 5 d_b, ที่มาของข้อมูล) แถบข้อมูลร่วมของทุกชุดอยู่ด้านบน"],
 ["Fig. 2", "แผนผังตัวอย่างวาดตามสัดส่วนจริง (ก้อน 150 mm, l_e = 140 mm) ใช้ลายแรเงามาตรฐาน: จุดสำหรับคอนกรีต เส้นเฉียงสำหรับเหล็ก ลายตาข่ายสำหรับสารยึดเหนี่ยว ตัดคำบรรยายที่ฝังบนภาพถ่าย"],
 ["Fig. 3", "3 panel ตาม RB6/9/12 กริดบาง เส้นหนา 1.1 pt สัญลักษณ์ตามชนิดสารยึดเหนี่ยว ใช้ error bar = ±1 SD"],
 ["Fig. 4", "(a)(b) ค่าเฉลี่ยพร้อม 95% CI (a: คำนวณจาก Table 3, b: รวมกลุ่มอย่างถูกต้องทางสถิติ) (c) แท่งมีลายแรเงาแยกตามสารยึดเหนี่ยว ใส่ตัวเลขร้อยละกำกับ"],
 ["Fig. 5", "τ และ σ_b ตาม d_b เส้นประ 235 MPa พร้อมหมายเหตุบนพื้นขาว σ_b คำนวณจาก τ ด้วย σ_b = 4 l_e τ / d_b"],
 ["Fig. 6", "ใช้ลายแรเงาแทนสีทึบ ตัวเลขในแท่งมีพื้นขาวรองเพื่ออ่านง่าย"],
 ["Fig. 7", "แกน log–log เส้นเฉียงตามเลขชี้กำลัง b ที่รายงาน แถบโปร่งแสงแสดงช่วง min–max ของทั้ง 5 ตัวอย่าง (จาก Table 7) และเส้นจุดแสดงกรณี P คงที่ที่ 20.8 kN"],
 ["Fig. 8", "แถบสีเทาแสดงค่าเฉลี่ย ± 1 SD รวมที่ f′c ≥ 14.7 MPa เส้นประ ACI 440.1R-06 พร้อมค่า 0.99 / 1.06 / 1.19 ค่าเฉลี่ย N รวมที่คำนวณได้ 1.62 / 1.52 / 1.11 ตรงกับ Section 3.8"],
 ["Fig. 9", "(a) dot-and-whisker แยกระบบ (b) GFRP เทียบกับ epoxy ที่ d_b และ f′c เดียวกัน เส้น 1:1 ป้ายเลข d_b ไม่ทับกัน"],
], [2.0, 15.2])
P("ข้อกำหนดที่ใช้ทุกรูป: ไม่มีข้อความ 'Fig. N' หรือชื่อรูปฝังในภาพ สีโทนหม่น อ่านได้เมื่อพิมพ์ขาวดำ ความละเอียด 420 dpi (กว้างไม่น้อยกว่า 2800 px) สคริปต์อยู่ที่ bond_paper/scripts")
H("4. การแก้ถ้อยคำ (เก่า → ใหม่)", 1)
log = json.load(open("manuscript/wording_log.json"))
T(["ตำแหน่ง", "เดิม", "ใหม่"], [[w, o, n] for w, o, n in log], [3.0, 6.6, 7.6])
H("5. ข้อเสนอแนะที่ยังไม่ได้แก้ (ขอให้ตัดสินใจ)", 1)
for t in [
 "ชื่อเรื่องใช้คำว่า 'Weak Concrete' ซึ่งไม่ได้นิยามและอิงกำลังอัดระบุ (nominal) ที่ reviewer ตั้งข้อสงสัย อาจเปลี่ยนเป็น 'Low-Strength Concrete' หรือเพิ่มช่วง 'Nominal 4.9–23.5 MPa' ในชื่อ ยังไม่ได้แก้เพราะเป็นการตัดสินใจของผู้เขียน",
 "Section 4.3 ระบุว่า RB12 ที่ 4.9 MPa บันทึกว่าเหล็กครากที่ 146 MPa ซึ่งขัดกับค่า catalogue ยังคงเป็นประเด็นที่ต้องตรวจกับสมุดบันทึกหรือไฟล์ดิบ",
 "ค่า CV ของ GFRP ที่ต่ำผิดปกติ (2–7%) ควรตรวจกับไฟล์ดิบ (อยู่ใน Table S5)",
]: d.add_paragraph(t, style="List Bullet")
d.save("manuscript/Bond_RB_GFRP_revision_log_TH.docx"); print("ok")
